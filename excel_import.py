import re
import secrets
import string
from io import BytesIO

from openpyxl import Workbook, load_workbook

from models import SchoolClass, Student, Teacher, User, db

PROTECTED_USERNAMES = {"admin", "superadmin", "super.admin", "root"}

STUDENT_ALIASES = {
    "full_name": ("fullname", "name", "studentname", "student", "fullnameofstudent"),
    "username": ("username", "user", "userid", "login", "loginid"),
    "email": ("email", "e-mail", "mail", "emailid", "emailaddress"),
    "admission_number": (
        "admissionnumber",
        "admissionno",
        "admission",
        "admno",
        "admissionn",
    ),
    "roll_number": ("rollnumber", "roll", "rollno", "rolln"),
    "grade": ("grade", "classgrade", "standard", "std"),
    "section": ("section", "sec", "division"),
    "class_label": ("class", "classsection", "gradeandsection"),
}

TEACHER_ALIASES = {
    "full_name": ("fullname", "name", "teachername", "teacher"),
    "username": ("username", "user", "userid", "login", "loginid"),
    "employee_id": (
        "employeeid",
        "employeeidnumber",
        "employeenumber",
        "employeeno",
        "empno",
        "empid",
        "staffid",
        "staffno",
        "staffnumber",
    ),
    "email": ("email", "e-mail", "mail", "emailid", "emailaddress"),
    "subject": ("subject", "department"),
    "grade": ("grade", "classgrade", "standard", "std"),
    "section": ("section", "sec", "division"),
    "class_label": ("class", "classsection", "gradeandsection"),
    "can_read": ("canread", "read", "readpermission"),
    "can_write": ("canwrite", "write", "writepermission"),
    "can_update": ("canupdate", "update", "updatepermission"),
    "can_deactivate": ("candeactivate", "deactivate", "deactivatestudents"),
}


def normalize_header(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def truthy(value):
    return cell_text(value).lower() in {"1", "yes", "y", "true", "on"}


def map_headers(headers, aliases):
    mapping = {}
    for index, header in enumerate(headers):
        key = normalize_header(header)
        if not key:
            continue
        for field, names in aliases.items():
            if field in mapping:
                continue
            if key == field or key in names:
                mapping[field] = index
                break
    return mapping


def read_excel_rows(file_storage):
    workbook = load_workbook(file_storage, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError("The Excel file is empty.")
    headers = [cell_text(item) for item in rows[0]]
    if not any(headers):
        raise ValueError("The first row must contain header names.")
    data = []
    for number, raw in enumerate(rows[1:], start=2):
        values = list(raw or [])
        if not any(cell_text(item) for item in values):
            continue
        data.append((number, values))
    return headers, data


def row_value(values, mapping, field):
    index = mapping.get(field)
    if index is None or index >= len(values):
        return ""
    return cell_text(values[index])


def random_password(length=8):
    alphabet = string.ascii_letters + string.digits
    return "".join(secrets.choice(alphabet) for _ in range(length))


def username_from_name(name, used):
    parts = re.findall(r"[a-z0-9]+", (name or "").lower())
    base = ".".join(parts)[:70] or "user"
    candidate = base
    suffix = 2
    while candidate in used or User.query.filter_by(username=candidate).first():
        candidate = f"{base}{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def unique_email(email, used, username):
    assigned, _reason = assign_email(email, used, username)
    return assigned


def normalize_person_name(name):
    return re.sub(r"\s+", " ", (name or "").strip().lower())


def assign_email(email, used, username):
    email = (email or "").strip().lower() or f"{username}@school.local"
    if email in used or User.query.filter(User.email == email).first():
        return None, f"email {email} is already registered."
    used.add(email)
    return email, None


def protected_identity_reason(username, email="", full_name=""):
    username = (username or "").strip().lower()
    email = (email or "").strip().lower()
    name = normalize_person_name(full_name)
    if username in PROTECTED_USERNAMES:
        return "that username is reserved for the super admin."
    for admin in User.query.filter_by(role="super_admin"):
        if username and admin.username.lower() == username:
            return "that username belongs to the super admin."
        if email and admin.email and admin.email.lower() == email:
            return "that email belongs to the super admin."
        if name and normalize_person_name(admin.full_name) == name:
            return "that name belongs to the super admin."
    return None


def normalize_id(value):
    return str(value or "").strip().lower()


def admission_taken(admission_number, exclude_student_id=None):
    needle = normalize_id(admission_number)
    if not needle:
        return False
    query = Student.query
    if exclude_student_id:
        query = query.filter(Student.id != exclude_student_id)
    for row in query:
        current = normalize_id(row.admission_number) or normalize_id(
            row.user.username if row.user else ""
        )
        if current == needle:
            return True
    return False


def employee_taken(employee_id, exclude_teacher_id=None):
    needle = normalize_id(employee_id)
    if not needle:
        return False
    query = Teacher.query
    if exclude_teacher_id:
        query = query.filter(Teacher.id != exclude_teacher_id)
    for row in query:
        current = normalize_id(row.employee_id) or normalize_id(
            row.user.username if row.user else ""
        )
        if current == needle:
            return True
    return False


def account_conflict(
    role,
    username,
    email="",
    full_name="",
    class_id=None,
    roll_number=None,
    admission_number=None,
    employee_id=None,
    exclude_user_id=None,
    exclude_student_id=None,
    exclude_teacher_id=None,
):
    username = (username or "").strip().lower()
    email = (email or "").strip().lower()
    protected = protected_identity_reason(username, email, full_name)
    if protected:
        return protected[0].upper() + protected[1:]
    if role == "student" and admission_number and admission_taken(
        admission_number, exclude_student_id
    ):
        return f"Admission number {admission_number} is already registered."
    if role == "teacher" and employee_id and employee_taken(
        employee_id, exclude_teacher_id
    ):
        return f"Employee ID {employee_id} is already registered."
    if username:
        taken = User.query.filter(User.username == username)
        if exclude_user_id:
            taken = taken.filter(User.id != exclude_user_id)
        if taken.first():
            return "That username is already used."
    if email:
        taken = User.query.filter(User.email == email)
        if exclude_user_id:
            taken = taken.filter(User.id != exclude_user_id)
        if taken.first():
            return "That email is already registered."
    return None


def grade_and_section(values, mapping):
    grade = row_value(values, mapping, "grade")
    section = row_value(values, mapping, "section").upper()
    label = row_value(values, mapping, "class_label")
    if (not grade or not section) and label:
        parts = [part for part in re.split(r"[-/\s]+", label.strip()) if part]
        if not grade and parts:
            grade = parts[0]
        if not section and len(parts) > 1:
            section = parts[1].upper()
    return grade, section


def find_class(values, mapping, default_class_id=None):
    grade, section = grade_and_section(values, mapping)
    if grade and section:
        found = SchoolClass.query.filter(
            SchoolClass.grade == grade, SchoolClass.section == section
        ).first()
        if found:
            return found
    if default_class_id:
        return db.session.get(SchoolClass, default_class_id)
    return None


def find_or_create_class(values, mapping, default_class_id=None, cache=None):
    cache = cache if cache is not None else {}
    grade, section = grade_and_section(values, mapping)
    if not grade:
        return find_class(values, mapping, default_class_id), False
    if not section:
        section = "A"

    key = (str(grade), section)
    if key in cache:
        return cache[key], False

    found = SchoolClass.query.filter(
        SchoolClass.grade == str(grade), SchoolClass.section == section
    ).first()
    if found:
        cache[key] = found
        return found, False

    school_class = SchoolClass(
        name=f"Class {grade}-{section}",
        grade=str(grade),
        section=section,
    )
    db.session.add(school_class)
    db.session.flush()
    cache[key] = school_class
    return school_class, True


def parse_student_rows(file_storage, default_class_id=None):
    headers, rows = read_excel_rows(file_storage)
    mapping = map_headers(headers, STUDENT_ALIASES)
    used_usernames = set()
    used_emails = set()
    used_admissions = set()
    class_cache = {}
    created_classes = []
    ready = []
    errors = []

    if "full_name" not in mapping:
        raise ValueError("Add a Full name column.")
    if "admission_number" not in mapping and "username" not in mapping and "roll_number" not in mapping:
        raise ValueError("Add an Admission number column. That value becomes the student username.")

    for number, values in rows:
        full_name = row_value(values, mapping, "full_name")
        admission_number = (
            row_value(values, mapping, "admission_number")
            or row_value(values, mapping, "username")
            or row_value(values, mapping, "roll_number")
        )
        email = row_value(values, mapping, "email").lower()
        roll_number = row_value(values, mapping, "roll_number") or admission_number
        username = admission_number.lower()

        if not full_name:
            errors.append(f"Row {number}: name is required.")
            continue
        if not admission_number:
            errors.append(f"Row {number}: admission number is required. It becomes the username.")
            continue
        school_class, class_created = find_or_create_class(
            values, mapping, default_class_id, class_cache
        )
        if not school_class:
            errors.append(f"Row {number}: add Grade and Section so the class can be found or created.")
            continue
        admission_key = normalize_id(admission_number)
        if username in used_usernames or admission_key in used_admissions:
            errors.append(f"Row {number}: skipped duplicate admission number {admission_number}.")
            continue
        if email and email in used_emails:
            errors.append(f"Row {number}: skipped duplicate email {email}.")
            continue
        conflict = account_conflict(
            "student",
            username,
            email=email,
            full_name=full_name,
            class_id=school_class.id,
            roll_number=roll_number,
            admission_number=admission_number,
        )
        if conflict:
            errors.append(f"Row {number}: skipped. {conflict}")
            continue
        email, email_error = assign_email(email, used_emails, username)
        if email_error:
            errors.append(f"Row {number}: skipped. {email_error[0].upper() + email_error[1:]}")
            continue
        used_usernames.add(username)
        used_admissions.add(admission_key)
        if class_created:
            created_classes.append(f"{school_class.grade}-{school_class.section}")
        ready.append(
            {
                "row": number,
                "full_name": full_name,
                "username": username,
                "email": email,
                "password": random_password(),
                "admission_number": admission_number,
                "roll_number": roll_number,
                "class_id": school_class.id,
            }
        )
    return ready, errors, created_classes


def parse_teacher_rows(file_storage, default_class_id=None, default_rights=None):
    headers, rows = read_excel_rows(file_storage)
    mapping = map_headers(headers, TEACHER_ALIASES)
    used_usernames = set()
    used_emails = set()
    used_employee_ids = set()
    ready = []
    errors = []

    if "full_name" not in mapping and "username" not in mapping:
        raise ValueError("Add a Full name or Username column.")

    for number, values in rows:
        full_name = row_value(values, mapping, "full_name") or row_value(values, mapping, "username")
        employee_id = row_value(values, mapping, "employee_id") or row_value(values, mapping, "username")
        username = (row_value(values, mapping, "username") or employee_id).lower()
        email = row_value(values, mapping, "email").lower()
        subject = row_value(values, mapping, "subject")
        school_class = find_class(values, mapping, default_class_id)

        if not full_name:
            errors.append(f"Row {number}: name is required.")
            continue
        if not employee_id:
            errors.append(f"Row {number}: employee ID is required.")
            continue
        employee_key = normalize_id(employee_id)
        if employee_key in used_employee_ids:
            errors.append(f"Row {number}: skipped duplicate employee ID {employee_id}.")
            continue
        if username and username in used_usernames:
            errors.append(f"Row {number}: skipped duplicate username {username}.")
            continue
        if email and email in used_emails:
            errors.append(f"Row {number}: skipped duplicate email {email}.")
            continue
        if not username:
            username = username_from_name(full_name, used_usernames)
        conflict = account_conflict(
            "teacher",
            username,
            email=email,
            full_name=full_name,
            employee_id=employee_id,
        )
        if conflict:
            errors.append(f"Row {number}: skipped. {conflict}")
            continue
        email, email_error = assign_email(email, used_emails, username)
        if email_error:
            errors.append(f"Row {number}: skipped. {email_error[0].upper() + email_error[1:]}")
            continue
        used_usernames.add(username)
        used_employee_ids.add(employee_key)
        ready.append(
            {
                "row": number,
                "full_name": full_name,
                "username": username,
                "employee_id": employee_id,
                "email": email,
                "password": random_password(),
                "subject": subject,
                "class_id": school_class.id if school_class else None,
                "can_read": True,
                "can_write": False,
                "can_update": False,
                "can_deactivate": False,
            }
        )
    return ready, errors


def excel_template(kind):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = kind.title()
    if kind == "students":
        sheet.append(
            ["Full name", "Admission number", "Email", "Roll number", "Grade", "Section"]
        )
        sheet.append(["Aarav Patel", "10A10", "aarav.patel@school.local", "10A10", "10", "A"])
    else:
        sheet.append(["Full name", "Employee ID", "Username", "Email", "Subject", "Grade", "Section"])
        sheet.append(["Anita Sharma", "T001", "anita.sharma", "anita.sharma@school.local", "Science", "10", "A"])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def credentials_workbook(kind, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Temporary passwords"
    detail = "Admission number" if kind == "students" else "Employee ID"
    include_registered = any(row.get("registered") for row in rows)
    headers = ["Name", "Username", "Temporary password", detail]
    if include_registered:
        headers.append("Registered")
    sheet.append(headers)
    for row in rows:
        password = row.get("password") or (
            "Changed by user" if row.get("status") == "changed" else "Not stored"
        )
        values = [row.get("name"), row.get("username"), password, row.get("detail")]
        if include_registered:
            values.append(row.get("registered") or "")
        sheet.append(values)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
