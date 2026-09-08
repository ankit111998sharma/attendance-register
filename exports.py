from collections import defaultdict
from datetime import date
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from mail_utils import get_setting
from models import Attendance


HEADERS = [
    "Roll",
    "Name",
    "Admission number",
    "Username",
    "Email",
    "Face photos",
    "Access",
    "Attendance",
]


def school_title(custom=""):
    custom = (custom or "").strip()
    if custom:
        return custom
    return get_setting("school_name", "").strip() or "Attendance Register"


def safe_filename(value):
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", (value or "").strip())
    return cleaned.strip("._") or "students"


def attendance_summary_map(student_ids):
    stats = {
        student_id: {"present": 0, "late": 0, "absent": 0, "total": 0, "percent": 0}
        for student_id in student_ids
    }
    if not student_ids:
        return stats
    for row in Attendance.query.filter(Attendance.student_id.in_(student_ids)).all():
        item = stats[row.student_id]
        item["total"] += 1
        if row.status in item:
            item[row.status] += 1
    for item in stats.values():
        if item["total"]:
            item["percent"] = round((item["present"] + item["late"]) / item["total"] * 100)
    return stats


def student_record(student, stats):
    school_class = student.school_class
    user = student.user
    att = stats.get(student.id) or {}
    return {
        "roll": student.roll_number or "",
        "name": user.full_name,
        "admission": student.admission_number or user.username,
        "username": user.username,
        "email": user.email or "",
        "grade": school_class.grade,
        "section": school_class.section,
        "class_name": school_class.name,
        "class_label": f"{school_class.grade}-{school_class.section} · {school_class.name}",
        "registered": user.created_at.strftime("%d %b %Y") if user.created_at else "",
        "face": f"{student.face_photo_count or 0}{' ready' if student.face_enrolled else ''}".strip(),
        "access": "Active" if user.is_active else "Deactivated",
        "attendance": (
            f"{att.get('percent', 0)}% ({att.get('present', 0)}P {att.get('late', 0)}L {att.get('absent', 0)}A)"
            if att.get("total")
            else "No marks"
        ),
    }


def grouped_student_rows(students):
    stats = attendance_summary_map([student.id for student in students])
    groups = defaultdict(list)
    for student in students:
        record = student_record(student, stats)
        groups[(record["grade"], record["section"])].append(record)
    ordered = []
    for key in sorted(groups, key=lambda item: (len(item[0]), item[0], item[1])):
        rows = sorted(groups[key], key=lambda item: (item["roll"], item["name"]))
        ordered.append((f"{key[0]}-{key[1]}", rows))
    return ordered


def filter_caption(filters, fallback="All records"):
    parts = []
    if filters.get("grade"):
        parts.append(f"Grade {filters['grade']}")
    if filters.get("section"):
        parts.append(f"Section {filters['section']}")
    if filters.get("search"):
        parts.append(f"Search “{filters['search']}”")
    if filters.get("access"):
        parts.append(filters["access"].title())
    if filters.get("face"):
        parts.append(f"Face {filters['face']}")
    if filters.get("reports") == "yes":
        parts.append("Reports access")
    elif filters.get("reports") == "no":
        parts.append("No reports access")
    if filters.get("date_label"):
        parts.append(filters["date_label"])
    return " · ".join(parts) or fallback


def export_filename(filters, fmt, prefix="students"):
    stamp = date.today().isoformat()
    if filters.get("class_label"):
        label = safe_filename(filters["class_label"].split("·")[0])
    elif filters.get("grade") or filters.get("section"):
        label = safe_filename(f"{filters.get('grade') or 'all'}-{filters.get('section') or 'all'}")
    else:
        label = "all"
    return f"{prefix}_{label}_{stamp}.{fmt}"


def record_values(row):
    return [
        row["roll"],
        row["name"],
        row["admission"],
        row["username"],
        row["email"],
        row["face"],
        row["access"],
        row["attendance"],
    ]


def build_xlsx(groups, heading, caption):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    navy = "1D3552"
    paper = "F4E8CC"
    thin = Border(
        left=Side(style="thin", color="C4A574"),
        right=Side(style="thin", color="C4A574"),
        top=Side(style="thin", color="C4A574"),
        bottom=Side(style="thin", color="C4A574"),
    )
    title_font = Font(name="Calibri", size=18, bold=True, color=navy)
    subtitle_font = Font(name="Calibri", size=11, color="6D6254")
    header_font = Font(name="Calibri", size=10, bold=True, color="2A1D08")
    group_font = Font(name="Calibri", size=12, bold=True, color=navy)

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    sheet["A1"] = heading
    sheet["A1"].font = title_font
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 26

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
    sheet["A2"] = f"Student register · {caption} · {date.today().strftime('%d %B %Y')}"
    sheet["A2"].font = subtitle_font

    row_index = 4
    for group_name, rows in groups:
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=len(HEADERS))
        cell = sheet.cell(row=row_index, column=1, value=f"{group_name} · {len(rows)} students")
        cell.font = group_font
        row_index += 1
        for col, header in enumerate(HEADERS, start=1):
            head = sheet.cell(row=row_index, column=col, value=header)
            head.font = header_font
            head.fill = PatternFill("solid", fgColor=paper)
            head.border = thin
        row_index += 1
        for record in rows:
            for col, value in enumerate(record_values(record), start=1):
                cell = sheet.cell(row=row_index, column=col, value=value)
                cell.border = thin
                cell.alignment = Alignment(vertical="center")
            row_index += 1
        row_index += 1

    widths = [10, 28, 18, 16, 28, 14, 14, 24]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddHeader.left.text = heading
    sheet.oddFooter.right.text = "Page &P of &N"

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def build_pdf(groups, heading, caption):
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=heading,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontName="Times-Bold",
        fontSize=18,
        textColor=colors.HexColor("#1d3552"),
        spaceAfter=2,
    )
    sub_style = ParagraphStyle(
        "ReportSub",
        parent=styles["Normal"],
        fontName="Times-Roman",
        fontSize=10,
        textColor=colors.HexColor("#6d6254"),
        spaceAfter=10,
    )
    group_style = ParagraphStyle(
        "ReportGroup",
        parent=styles["Heading2"],
        fontName="Times-Bold",
        fontSize=12,
        textColor=colors.HexColor("#1d3552"),
        spaceBefore=8,
        spaceAfter=4,
    )
    cell_style = ParagraphStyle("Cell", fontName="Times-Roman", fontSize=8, leading=10)
    head_style = ParagraphStyle("Head", fontName="Times-Bold", fontSize=8, leading=10)

    story = [
        Paragraph(heading, title_style),
        Paragraph(f"Student register · {caption} · {date.today().strftime('%d %B %Y')}", sub_style),
    ]
    for group_name, rows in groups:
        story.append(Paragraph(f"{group_name} · {len(rows)} students", group_style))
        data = [[Paragraph(header, head_style) for header in HEADERS]]
        for record in rows:
            data.append([Paragraph(str(value or "—"), cell_style) for value in record_values(record)])
        table = Table(data, colWidths=[22 * mm, 45 * mm, 32 * mm, 28 * mm, 48 * mm, 24 * mm, 24 * mm, 40 * mm], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f4e8cc")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#2a1d08")),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#c4a574")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fff8ea")]),
                ]
            )
        )
        story.append(table)
        story.append(Spacer(1, 8))
    document.build(story)
    buffer.seek(0)
    return buffer


def build_docx(groups, heading, caption):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run(heading)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1D, 0x35, 0x52)

    sub = document.add_paragraph()
    run = sub.add_run(f"Student register · {caption} · {date.today().strftime('%d %B %Y')}")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x6D, 0x62, 0x54)

    for group_name, rows in groups:
        heading_p = document.add_paragraph()
        run = heading_p.add_run(f"{group_name} · {len(rows)} students")
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x1D, 0x35, 0x52)
        table = document.add_table(rows=1, cols=len(HEADERS))
        table.style = "Table Grid"
        for index, header in enumerate(HEADERS):
            table.rows[0].cells[index].text = header
        for record in rows:
            cells = table.add_row().cells
            for index, value in enumerate(record_values(record)):
                cells[index].text = str(value or "—")
        document.add_paragraph()

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


MIME_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def build_student_export(students, heading, filters, fmt):
    groups = grouped_student_rows(students)
    caption = filter_caption(filters, "All grades and sections")
    if fmt == "pdf":
        payload = build_pdf(groups, heading, caption)
    elif fmt == "docx":
        payload = build_docx(groups, heading, caption)
    else:
        fmt = "xlsx"
        payload = build_xlsx(groups, heading, caption)
    return payload, export_filename(filters, fmt), MIME_TYPES[fmt]


TEACHER_HEADERS = [
    "Name",
    "Employee ID",
    "Username",
    "Email",
    "Subject",
    "Grade",
    "Section",
    "Rights",
    "Access",
    "Reports",
    "Registered",
]


def teacher_rights_text(perm):
    if not perm:
        return "None"
    labels = []
    if perm.can_read:
        labels.append("Read")
    if perm.can_write:
        labels.append("Write")
    if perm.can_update:
        labels.append("Update")
    if perm.can_deactivate:
        labels.append("Deactivate")
    return ", ".join(labels) or "None"


def teacher_record_values(row):
    return [
        row["name"],
        row["employee_id"],
        row["username"],
        row["email"],
        row["subject"],
        row["grade"],
        row["section"],
        row["rights"],
        row["access"],
        row["reports"],
        row["registered"],
    ]


def grouped_teacher_rows(teachers):
    groups = defaultdict(list)
    for teacher in teachers:
        user = teacher.user
        base = {
            "name": user.full_name,
            "employee_id": teacher.employee_id or "",
            "username": user.username,
            "email": user.email or "",
            "subject": teacher.subject or "",
            "access": "Active" if user.is_active else "Deactivated",
            "reports": "Yes" if teacher.can_reports else "No",
            "registered": user.created_at.strftime("%d %b %Y") if user.created_at else "",
        }
        if teacher.permissions:
            for perm in teacher.permissions:
                groups[(perm.school_class.grade, perm.school_class.section)].append(
                    {
                        **base,
                        "grade": perm.school_class.grade,
                        "section": perm.school_class.section,
                        "rights": teacher_rights_text(perm),
                    }
                )
        else:
            groups[("", "")].append(
                {**base, "grade": "—", "section": "—", "rights": "None"}
            )
    ordered = []
    for key in sorted(groups, key=lambda item: (item[0] == "", len(item[0]), item[0], item[1])):
        label = f"{key[0]}-{key[1]}" if key[0] else "No grade assigned"
        rows = sorted(groups[key], key=lambda item: (item["name"], item["employee_id"]))
        ordered.append((label, rows))
    return ordered


def build_teacher_xlsx(groups, heading, caption):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Teachers"
    navy = "1D3552"
    paper = "F4E8CC"
    thin = Border(
        left=Side(style="thin", color="C4A574"),
        right=Side(style="thin", color="C4A574"),
        top=Side(style="thin", color="C4A574"),
        bottom=Side(style="thin", color="C4A574"),
    )
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(TEACHER_HEADERS))
    sheet["A1"] = heading
    sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color=navy)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(TEACHER_HEADERS))
    sheet["A2"] = f"Teacher register · {caption} · {date.today().strftime('%d %B %Y')}"
    sheet["A2"].font = Font(name="Calibri", size=11, color="6D6254")
    row_index = 4
    for group_name, rows in groups:
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=len(TEACHER_HEADERS))
        sheet.cell(row=row_index, column=1, value=f"{group_name} · {len(rows)} rows").font = Font(
            name="Calibri", size=12, bold=True, color=navy
        )
        row_index += 1
        for col, header in enumerate(TEACHER_HEADERS, start=1):
            head = sheet.cell(row=row_index, column=col, value=header)
            head.font = Font(name="Calibri", size=10, bold=True, color="2A1D08")
            head.fill = PatternFill("solid", fgColor=paper)
            head.border = thin
        row_index += 1
        for record in rows:
            for col, value in enumerate(teacher_record_values(record), start=1):
                cell = sheet.cell(row=row_index, column=col, value=value)
                cell.border = thin
            row_index += 1
        row_index += 1
    for index, width in enumerate([22, 14, 16, 26, 16, 10, 10, 22, 14, 12, 14], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.page_setup.orientation = "landscape"
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def build_teacher_pdf(groups, heading, caption):
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=heading,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TTitle", parent=styles["Heading1"], fontName="Times-Bold", fontSize=18,
        textColor=colors.HexColor("#1d3552"), spaceAfter=2,
    )
    sub_style = ParagraphStyle(
        "TSub", parent=styles["Normal"], fontName="Times-Roman", fontSize=10,
        textColor=colors.HexColor("#6d6254"), spaceAfter=10,
    )
    group_style = ParagraphStyle(
        "TGroup", parent=styles["Heading2"], fontName="Times-Bold", fontSize=12,
        textColor=colors.HexColor("#1d3552"), spaceBefore=8, spaceAfter=4,
    )
    cell_style = ParagraphStyle("TCell", fontName="Times-Roman", fontSize=7, leading=9)
    head_style = ParagraphStyle("THead", fontName="Times-Bold", fontSize=7, leading=9)
    story = [
        Paragraph(heading, title_style),
        Paragraph(f"Teacher register · {caption} · {date.today().strftime('%d %B %Y')}", sub_style),
    ]
    widths = [28 * mm, 20 * mm, 22 * mm, 36 * mm, 22 * mm, 14 * mm, 16 * mm, 32 * mm, 18 * mm, 16 * mm, 22 * mm]
    for group_name, rows in groups:
        story.append(Paragraph(f"{group_name} · {len(rows)} rows", group_style))
        data = [[Paragraph(header, head_style) for header in TEACHER_HEADERS]]
        for record in rows:
            data.append([Paragraph(str(value or "—"), cell_style) for value in teacher_record_values(record)])
        table = Table(data, colWidths=widths, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f4e8cc")),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#c4a574")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fff8ea")]),
                ]
            )
        )
        story.append(table)
    document.build(story)
    buffer.seek(0)
    return buffer


def build_teacher_docx(groups, heading, caption):
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    title = document.add_paragraph()
    run = title.add_run(heading)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1D, 0x35, 0x52)
    sub = document.add_paragraph()
    run = sub.add_run(f"Teacher register · {caption} · {date.today().strftime('%d %B %Y')}")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x6D, 0x62, 0x54)
    for group_name, rows in groups:
        heading_p = document.add_paragraph()
        run = heading_p.add_run(f"{group_name} · {len(rows)} rows")
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x1D, 0x35, 0x52)
        table = document.add_table(rows=1, cols=len(TEACHER_HEADERS))
        table.style = "Table Grid"
        for index, header in enumerate(TEACHER_HEADERS):
            table.rows[0].cells[index].text = header
        for record in rows:
            cells = table.add_row().cells
            for index, value in enumerate(teacher_record_values(record)):
                cells[index].text = str(value or "—")
        document.add_paragraph()
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def build_teacher_export(teachers, heading, filters, fmt):
    groups = grouped_teacher_rows(teachers)
    caption = filter_caption(filters, "All teachers")
    if fmt == "pdf":
        payload = build_teacher_pdf(groups, heading, caption)
    elif fmt == "docx":
        payload = build_teacher_docx(groups, heading, caption)
    else:
        fmt = "xlsx"
        payload = build_teacher_xlsx(groups, heading, caption)
    return payload, export_filename(filters, fmt, "teachers"), MIME_TYPES[fmt]
